#!/usr/bin/env python3
"""Generate a PRIVATE Supabase seed from the complete BME Activity Tracker XLSX.

Usage:
    python scripts/generate_full_seed.py "BME - Jan 26 - Activity Tracker - Final.xlsx"

Output:
    supabase/seed.local.sql

The output contains real faculty/course responsibility data. It is intentionally written
to the existing git-ignored seed.local.sql path. Never commit that generated SQL.

Dependency:
    pip install openpyxl
"""
from __future__ import annotations

import datetime as dt
import pathlib
import re
import sys
from collections import defaultdict
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / "supabase" / "seed.local.sql"

TERM_NAME = "January 2026"
TERM_LEGACY = "2026"
TOTAL_TEACHING_WEEKS = 14
LATEST_COMPLETED_WEEK = 9


def q(value: Any) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


def norm_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def norm_code(value: Any) -> str:
    return norm_text(value)


def norm_assessment(value: Any) -> str:
    text = norm_text(value)
    fixes = {
        "Conitnuous": "Continuous",
        "Continous": "Continuous",
        "Submissons": "Submissions",
    }
    for old, new in fixes.items():
        text = text.replace(old, new)
    return text


def iso_date(value: Any, epoch) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value, epoch=epoch)
            if isinstance(converted, dt.datetime):
                converted = converted.date()
            return converted.isoformat()
        except Exception:
            pass
    raw = norm_text(value)
    raw = raw.replace("Auguest", "August")
    raw = raw.replace("Septemeber", "September")
    fmts = (
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d-%b-%Y",
        "%A, %B %d, %Y", "%B %d, %Y", "%B %d,%Y",
    )
    for fmt in fmts:
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized date value: {value!r}")


def ct_status(value: Any) -> str:
    return {
        "CT Scheduled": "scheduled",
        "CT Taken": "ct_taken",
        "Scripts Under Examination": "scripts_under_examination",
        "Scripts Checked": "scripts_checked",
        "Marks Published": "marks_published",
    }.get(norm_text(value), "not_planned")


def sessional_status(value: Any) -> str:
    return {
        "Scheduled": "scheduled",
        "Evaluation Taken": "evaluation_taken",
        "Marking Complete": "marking_complete",
        "Marks Uploaded in Excel": "marks_uploaded",
    }.get(norm_text(value), "scheduled")


def outline_status(value: Any) -> str:
    return {
        "Under Preparation": "under_preparation",
        "Prepared": "prepared",
        "Shared with Students": "shared_with_students",
    }.get(norm_text(value), "not_started")


def feedback_status(value: Any) -> str:
    return {
        "Students Reminded": "in_progress",
        "Feedbacks Collected": "completed",
    }.get(norm_text(value), "not_started")


def car_status(value: Any) -> str:
    return {
        "Under Preparation": "in_progress",
        "Prepared": "in_progress",
        "Uploaded in Drive": "uploaded",
    }.get(norm_text(value), "not_started")


def theory_status(stage: str, value: Any) -> str:
    text = norm_text(value)
    if not text:
        return "not_started"
    complete = {
        "question": {"Submitted"},
        "moderation": {"Finalized"},
        "exam": {"Forwarded to Script Scrutinizer"},
        "scrutiny": {"Forwarded to Gradesheet Preparer"},
        "grade_prep": {"Forwarded to Gradesheet Scrutinizer"},
        "grade_scrutiny": {"Forwarded to HOD"},
    }
    submitted = {
        "grade_prep": {"Uploaded in BIIS"},
        "grade_scrutiny": {"Submitted"},
    }
    if text in complete.get(stage, set()):
        return "completed"
    if text in submitted.get(stage, set()):
        return "submitted"
    return "in_progress"


def sql_course_id(code: str) -> str:
    return f"(select id from public.courses where course_code={q(code)} limit 1)"


def sql_faculty_id(name: str | None) -> str:
    return "null" if not name else f"(select id from public.faculty where full_name={q(name)} limit 1)"


def first_nonblank(ws, col: int, start=2, end=None):
    end = end or ws.max_row
    for r in range(start, end + 1):
        v = ws.cell(r, col).value
        if v not in (None, ""):
            return v
    return None


def collect_course_lists(wb):
    courses: dict[str, dict] = {}
    course_teachers: list[dict] = []
    names: set[str] = set()

    # Theory courses
    ws = wb["01_Theory Course List"]
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 3).value)
        if not code:
            continue
        title = norm_text(ws.cell(r, 4).value)
        contact = ws.cell(r, 5).value
        section = norm_text(ws.cell(r, 6).value) or "X"
        courses[code] = dict(code=code, title=title, contact=contact, course_type="theory", subtype=None)
        for order, col in enumerate((7, 8), start=1):
            name = norm_text(ws.cell(r, col).value)
            if name:
                names.add(name)
                course_teachers.append(dict(code=code, section=section, name=name, order=order))

    # Sessional courses: course code/title cells are merged, so carry forward previous course.
    ws = wb["02_Sessional Course List"]
    current_code = None
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 3).value) or current_code
        if not code:
            continue
        if ws.cell(r, 3).value:
            current_code = code
            title = norm_text(ws.cell(r, 4).value)
            contact = ws.cell(r, 5).value
            subtype = "THESIS" if "thesis" in title.lower() else ("DESIGN" if "design" in title.lower() else "LABORATORY")
            courses[code] = dict(code=code, title=title, contact=contact, course_type="sessional", subtype=subtype)
        section = norm_text(ws.cell(r, 6).value) or "X"
        for order, col in enumerate(range(7, 12), start=1):
            name = norm_text(ws.cell(r, col).value)
            if name:
                names.add(name)
                course_teachers.append(dict(code=code, section=section, name=name, order=order))
    return courses, course_teachers, names


def collect_theory(wb, names):
    ws = wb["08_Theory Grade Tracker"]
    rows = []
    for r in range(3, ws.max_row + 1):
        code = norm_code(ws.cell(r, 2).value)
        if not code:
            continue
        vals = {
            "code": code,
            "preparer_a": norm_text(ws.cell(r, 4).value),
            "prep_a_status": theory_status("question", ws.cell(r, 5).value),
            "preparer_b": norm_text(ws.cell(r, 6).value),
            "prep_b_status": theory_status("question", ws.cell(r, 7).value),
            "moderator_1": norm_text(ws.cell(r, 9).value),
            "mod1_status": theory_status("moderation", ws.cell(r, 10).value),
            "moderator_2": norm_text(ws.cell(r, 11).value),
            "mod2_status": theory_status("moderation", ws.cell(r, 12).value),
            "exam_a_status": theory_status("exam", ws.cell(r, 14).value),
            "exam_b_status": theory_status("exam", ws.cell(r, 16).value),
            "scrutinizer": norm_text(ws.cell(r, 17).value),
            "scrutiny_status": theory_status("scrutiny", ws.cell(r, 18).value),
            "grade_prep_status": theory_status("grade_prep", ws.cell(r, 20).value),
            "grade_scrutiny_status": theory_status("grade_scrutiny", ws.cell(r, 22).value),
        }
        for k in ("preparer_a", "preparer_b", "moderator_1", "moderator_2", "scrutinizer"):
            if vals[k]: names.add(vals[k])
        rows.append(vals)
    return rows


def collect_cts(wb):
    ws = wb["04_Class Test Tracker"]
    out = []
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 2).value)
        if not code:
            continue
        notes = norm_text(ws.cell(r, 16).value)
        for n in range(1, 6):
            date_col = 6 + (n - 1) * 2
            status_col = date_col + 1
            raw_date, raw_status = ws.cell(r, date_col).value, ws.cell(r, status_col).value
            if raw_date in (None, "") and raw_status in (None, ""):
                continue
            out.append(dict(
                code=code, number=n, title=f"CT {n:02d}",
                date=iso_date(raw_date, wb.epoch) if raw_date not in (None, "") else None,
                status=ct_status(raw_status), notes=notes,
            ))
    return out


def collect_outlines(wb):
    ws = wb["03_Course Outline Tracker"]
    out = {}
    current_code = None
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 3).value) or current_code
        if not code: continue
        if ws.cell(r, 3).value: current_code = code
        if code not in out:
            out[code] = outline_status(ws.cell(r, 11).value)
    deadline_raw = first_nonblank(ws, 12)
    return out, iso_date(deadline_raw, wb.epoch) if deadline_raw not in (None, "") else None


def collect_feedback(wb):
    ws = wb["06_Feedback Tracker"]
    out = {}
    current_code = None
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 3).value) or current_code
        if not code: continue
        if ws.cell(r, 3).value: current_code = code
        if code not in out:
            out[code] = feedback_status(ws.cell(r, 11).value)
    start = first_nonblank(ws, 12)
    deadline = first_nonblank(ws, 13)
    return out, (iso_date(start, wb.epoch) if start not in (None, "") else None), (iso_date(deadline, wb.epoch) if deadline not in (None, "") else None)


def collect_car(wb):
    ws = wb["09_CF_CAR Tracker"]
    out = {}
    current_code = None
    for r in range(2, ws.max_row + 1):
        code = norm_code(ws.cell(r, 3).value) or current_code
        if not code: continue
        if ws.cell(r, 3).value: current_code = code
        if code not in out:
            out[code] = car_status(ws.cell(r, 11).value)
    deadline = first_nonblank(ws, 12)
    return out, iso_date(deadline, wb.epoch) if deadline not in (None, "") else None


def collect_sessional_grades(wb, names):
    ws = wb["07_Sessional Grade Tracker"]
    out=[]
    for r in range(2, ws.max_row + 1):
        code=norm_code(ws.cell(r,2).value)
        if not code: continue
        prep=norm_text(ws.cell(r,6).value); scr=norm_text(ws.cell(r,8).value)
        if prep: names.add(prep)
        if scr: names.add(scr)
        out.append(dict(
            code=code, preparer=prep, scrutinizer=scr,
            prep_status=theory_status("grade_prep",ws.cell(r,7).value),
            scrutiny_status=theory_status("grade_scrutiny",ws.cell(r,9).value),
        ))
    deadline=first_nonblank(ws,10)
    return out, iso_date(deadline, wb.epoch) if deadline not in (None,"") else None


def collect_sessional_assessments(wb):
    ws=wb["05_Sessional Assessment Tracker"]
    records=[]; issues=[]; current_code=None
    for r in range(2,ws.max_row+1):
        code=norm_code(ws.cell(r,3).value) or current_code
        if not code: continue
        if ws.cell(r,3).value: current_code=code
        # Values normally live on first row of merged A1/A2 pair. Skip totally empty triplets.
        for slot, col in enumerate(range(7, ws.max_column+1, 3), start=1):
            raw_type=ws.cell(r,col).value
            raw_date=ws.cell(r,col+1).value if col+1<=ws.max_column else None
            raw_status=ws.cell(r,col+2).value if col+2<=ws.max_column else None
            if raw_type in (None,"") and raw_date in (None,"") and raw_status in (None,""):
                continue
            name=norm_assessment(raw_type)
            if not name:
                name=f"Unclassified Assessment {slot:02d}"
                issues.append(("05_Sessional Assessment Tracker",r,"missing_assessment_type","warning",f"{code}: evaluation slot {slot} has date/status but no evaluation type; preserved as {name}."))
            date=None
            if raw_date not in (None,""):
                try: date=iso_date(raw_date,wb.epoch)
                except ValueError as ex:
                    issues.append(("05_Sessional Assessment Tracker",r,"invalid_date","warning",f"{code} {name}: {ex}"))
            if name and not date:
                issues.append(("05_Sessional Assessment Tracker",r,"incomplete_assessment","info",f"{code}: {name} has no evaluation date in the workbook."))
            m=re.match(r"^(.*?)(?:\s+(\d{1,2}))$",name)
            if m:
                base=m.group(1).strip(); seq=int(m.group(2))
            else:
                base=name; seq=1
            records.append(dict(code=code,base=base,seq=seq,title=name,date=date,status=sessional_status(raw_status),source_row=r,slot=slot))
    return records,issues


def build_sql(path: pathlib.Path) -> str:
    wb=load_workbook(path,data_only=False)
    courses,teachers,names=collect_course_lists(wb)
    theory=collect_theory(wb,names)
    cts=collect_cts(wb)
    outlines,outline_deadline=collect_outlines(wb)
    feedback,feedback_start,feedback_deadline=collect_feedback(wb)
    car,car_deadline=collect_car(wb)
    sess_grades,sess_grade_deadline=collect_sessional_grades(wb,names)
    assessments,issues=collect_sessional_assessments(wb)

    # Existing theory sheet can contain courses before/after list order; keep course list as authority.
    out=[]
    out.append("-- GENERATED FROM THE FULL PRIVATE BME ACTIVITY TRACKER XLSX. DO NOT COMMIT.\n")
    out.append("-- Run latest schema.sql + activity_tracker_v3_migration.sql BEFORE this seed.\n\n")
    out.append("begin;\n\n")
    out.append(f"insert into public.academic_terms(name,total_teaching_weeks,latest_completed_week,calendar_status,active) values ({q(TERM_NAME)},{TOTAL_TEACHING_WEEKS},{LATEST_COMPLETED_WEEK},'active',true) on conflict do nothing;\n")
    out.append(f"update public.academic_terms set active=false where name<>{q(TERM_NAME)} and active=true;\n")
    out.append(f"update public.academic_terms set active=true,total_teaching_weeks={TOTAL_TEACHING_WEEKS},latest_completed_week={LATEST_COMPLETED_WEEK} where name={q(TERM_NAME)};\n\n")

    for name in sorted(names):
        out.append(f"insert into public.faculty(full_name) values ({q(name)}) on conflict (full_name) do nothing;\n")
    out.append("\n")

    # Draft first so theory publish validation does not fire before assignments exist.
    for c in courses.values():
        contact="null" if c['contact'] in (None,'') else str(float(c['contact']))
        subtype=q(c['subtype']) if c['subtype'] else "null"
        out.append(
            "insert into public.courses(course_code,title,term,term_id,course_type,sessional_subtype,contact_hours,status,active) "
            f"values ({q(c['code'])},{q(c['title'])},{q(TERM_LEGACY)},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(c['course_type'])},{subtype},{contact},'draft',true) "
            "on conflict (course_code) do update set title=excluded.title,term=excluded.term,term_id=excluded.term_id,course_type=excluded.course_type,sessional_subtype=excluded.sessional_subtype,contact_hours=excluded.contact_hours,active=true;\n"
        )
    out.append("\n")

    # Rebuild course sections/teacher links for this tracker source.
    course_codes=list(courses)
    code_list=",".join(q(c) for c in course_codes)
    out.append(f"delete from public.course_faculty where course_id in (select id from public.courses where course_code in ({code_list}));\n")
    out.append(f"delete from public.course_sections where course_id in (select id from public.courses where course_code in ({code_list}));\n")
    sections=sorted({(t['code'],t['section']) for t in teachers})
    for code,section in sections:
        out.append(f"insert into public.course_sections(course_id,section_code) values ({sql_course_id(code)},{q(section)}) on conflict (course_id,section_code) do nothing;\n")
    for t in teachers:
        out.append(
            "insert into public.course_faculty(course_id,section_id,faculty_id,role,section,display_order) values ("
            f"{sql_course_id(t['code'])},(select id from public.course_sections where course_id={sql_course_id(t['code'])} and section_code={q(t['section'])} limit 1),"
            f"{sql_faculty_id(t['name'])},'course_teacher',{q(t['section'])},{t['order']});\n"
        )
    out.append("\n")

    # Theory role assignments and coarse workflow status rows.
    for x in theory:
        out.append(
            "insert into public.primary_assignments(course_id,preparer_a,preparer_b,moderator_1,moderator_2,scrutinizer) values ("
            f"{sql_course_id(x['code'])},{sql_faculty_id(x['preparer_a'])},{sql_faculty_id(x['preparer_b'])},{sql_faculty_id(x['moderator_1'])},{sql_faculty_id(x['moderator_2'])},{sql_faculty_id(x['scrutinizer'])}) "
            "on conflict (course_id) do update set preparer_a=excluded.preparer_a,preparer_b=excluded.preparer_b,moderator_1=excluded.moderator_1,moderator_2=excluded.moderator_2,scrutinizer=excluded.scrutinizer;\n"
        )
        task_statuses={
            'question_prep_a':x['prep_a_status'],'question_prep_b':x['prep_b_status'],
            'moderation_1':x['mod1_status'],'moderation_2':x['mod2_status'],
            'examination_a':x['exam_a_status'],'examination_b':x['exam_b_status'],
            'scrutiny':x['scrutiny_status'],'gradesheet_prep':x['grade_prep_status'],
            'gradesheet_scrutiny':x['grade_scrutiny_status'],
        }
        # Only seed non-empty workbook progress; not_started is harmless but verbose.
        for key,status in task_statuses.items():
            if status!='not_started':
                out.append(f"insert into public.task_statuses(course_id,task_key,status) values ({sql_course_id(x['code'])},{q(key)},{q(status)}) on conflict (course_id,task_key) do update set status=excluded.status;\n")
    out.append("\n")

    # Shared term milestones.
    milestones=[
        ('COURSE_OUTLINE_SHARING',outline_deadline),('FEEDBACK_START',feedback_start),('FEEDBACK_DEADLINE',feedback_deadline),
        ('SESSIONAL_GRADESHEET_DEADLINE',sess_grade_deadline),('CAR_DEADLINE',car_deadline),
    ]
    out.append(f"delete from public.term_milestones where term_id=(select id from public.academic_terms where name={q(TERM_NAME)} limit 1) and course_id is null;\n")
    for kind,date in milestones:
        if date:
            out.append(f"insert into public.term_milestones(term_id,milestone_type,milestone_date,status) values ((select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(kind)},{q(date)},'scheduled');\n")
    out.append("\n")

    # Course outlines, feedback, CAR.
    for code in courses:
        out.append(f"insert into public.course_outline_statuses(course_id,term_id,status) values ({sql_course_id(code)},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(outlines.get(code,'not_started'))}) on conflict (course_id) do update set term_id=excluded.term_id,status=excluded.status;\n")
        out.append(f"insert into public.feedback_statuses(course_id,term_id,status) values ({sql_course_id(code)},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(feedback.get(code,'not_started'))}) on conflict (course_id) do update set term_id=excluded.term_id,status=excluded.status;\n")
        out.append(f"insert into public.course_file_statuses(course_id,term_id,status) values ({sql_course_id(code)},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(car.get(code,'not_started'))}) on conflict (course_id) do update set term_id=excluded.term_id,status=excluded.status;\n")
    out.append("\n")

    # Replace CT records for tracker courses with workbook CT records.
    out.append(f"delete from public.class_tests where course_id in (select id from public.courses where course_code in ({code_list}));\n")
    for ct in cts:
        out.append(
            "insert into public.class_tests(course_id,term_id,ct_number,title,section,scheduled_date,status,notes) values ("
            f"{sql_course_id(ct['code'])},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{ct['number']},{q(ct['title'])},'X',{q(ct['date']) if ct['date'] else 'null'},{q(ct['status'])},{q(ct['notes'])});\n"
        )
    out.append("\n")

    # Sessional configuration and dynamic assessment engine.
    sessional_codes=[c for c,v in courses.items() if v['course_type']=='sessional']
    sess_list=",".join(q(c) for c in sessional_codes)
    out.append(f"delete from public.assessment_instances where course_id in (select id from public.courses where course_code in ({sess_list}));\n")
    out.append(f"delete from public.assessment_components where course_id in (select id from public.courses where course_code in ({sess_list}));\n")
    for code in sessional_codes:
        out.append(f"insert into public.sessional_course_configs(course_id,term_id,session_count,config_status,version,remarks) values ({sql_course_id(code)},(select id from public.academic_terms where name={q(TERM_NAME)} limit 1),0,'setup_required',1,'Imported from workbook; assigned course teachers should confirm session count and assessment structure.') on conflict (course_id) do update set term_id=excluded.term_id,remarks=excluded.remarks;\n")

    grouped=defaultdict(list)
    for a in assessments: grouped[(a['code'],a['base'])].append(a)
    component_order=defaultdict(int)
    for (code,base),recs in grouped.items():
        component_order[code]+=1
        qty=max(r['seq'] for r in recs)
        out.append(f"insert into public.assessment_components(course_id,name,quantity,marks_each,counted_quantity,scope,display_order,active,remarks) values ({sql_course_id(code)},{q(base)},{qty},0,{qty},'Entire Course',{component_order[code]},true,'Imported from Activity Tracker; marks were not stored in this tracker.') ;\n")
        for rec in sorted(recs,key=lambda x:x['seq']):
            out.append(
                "insert into public.assessment_instances(component_id,course_id,sequence_number,title,scope,scheduled_date,status,notes) values ("
                f"(select ac.id from public.assessment_components ac where ac.course_id={sql_course_id(code)} and ac.name={q(base)} order by ac.created_at desc limit 1),"
                f"{sql_course_id(code)},{rec['seq']},{q(rec['title'])},'Entire Course',{q(rec['date']) if rec['date'] else 'null'},{q(rec['status'])},'Imported from workbook evaluation slot {rec['slot']}.');\n"
            )
    out.append("\n")

    # Sessional gradesheet responsibility.
    for g in sess_grades:
        out.append(f"insert into public.sessional_grade_assignments(course_id,preparer,scrutinizer) values ({sql_course_id(g['code'])},{sql_faculty_id(g['preparer'])},{sql_faculty_id(g['scrutinizer'])}) on conflict (course_id) do update set preparer=excluded.preparer,scrutinizer=excluded.scrutinizer;\n")
        out.append(f"insert into public.sessional_grade_statuses(course_id,task_key,status,responsible_faculty) values ({sql_course_id(g['code'])},'sessional_gradesheet_prep',{q(g['prep_status'])},{sql_faculty_id(g['preparer'])}) on conflict (course_id,task_key) do update set status=excluded.status,responsible_faculty=excluded.responsible_faculty;\n")
        out.append(f"insert into public.sessional_grade_statuses(course_id,task_key,status,responsible_faculty) values ({sql_course_id(g['code'])},'sessional_gradesheet_scrutiny',{q(g['scrutiny_status'])},{sql_faculty_id(g['scrutinizer'])}) on conflict (course_id,task_key) do update set status=excluded.status,responsible_faculty=excluded.responsible_faculty;\n")
    out.append("\n")

    # Import warnings preserve partial/inconsistent workbook records for cleanup.
    out.append(f"delete from public.import_validation_issues where term_id=(select id from public.academic_terms where name={q(TERM_NAME)} limit 1);\n")
    for sheet,row,itype,severity,msg in issues:
        out.append(f"insert into public.import_validation_issues(term_id,source_sheet,source_row,issue_type,severity,message) values ((select id from public.academic_terms where name={q(TERM_NAME)} limit 1),{q(sheet)},{row},{q(itype)},{q(severity)},{q(msg)});\n")

    # Publish only after theory assignments have been inserted.
    out.append(f"\nupdate public.courses set status='published' where course_code in ({code_list});\n")
    out.append("\ncommit;\n")
    out.append(f"\n-- Summary: {len(courses)} courses, {len(names)} faculty, {len(cts)} CT records, {len(assessments)} sessional assessment records, {len(issues)} import warnings.\n")
    return ''.join(out)


def main():
    if len(sys.argv)!=2:
        raise SystemExit('Usage: python scripts/generate_full_seed.py "BME - Jan 26 - Activity Tracker - Final.xlsx"')
    source=pathlib.Path(sys.argv[1]).expanduser().resolve()
    if not source.exists(): raise SystemExit(f"Workbook not found: {source}")
    OUT.parent.mkdir(parents=True,exist_ok=True)
    OUT.write_text(build_sql(source),encoding='utf-8')
    print(f"Wrote private Supabase seed: {OUT}")
    print("Next: run latest schema.sql, then activity_tracker_v3_migration.sql, then this seed in Supabase SQL Editor.")

if __name__=='__main__':
    main()
